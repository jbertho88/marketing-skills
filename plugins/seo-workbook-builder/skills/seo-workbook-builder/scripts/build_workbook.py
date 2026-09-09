"""
SEO Workbook Builder
Takes a Screaming Frog internal URLs export + ranking data and produces
a formatted, keyword-mapped SEO workbook.

Usage:
    python build_workbook.py \
        --crawl <path_to_internal_urls.xlsx> \
        --ranking <path_to_ranking_data.csv_or_xlsx> \
        --output <output_path.xlsx> \
        [--site-name "My Site"]
"""

import argparse
import json
import re
import sys
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ============================================================
# STYLES
# ============================================================
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill('solid', fgColor='1B2A4A')
DATA_FONT = Font(name='Arial', size=10)
LINK_FONT = Font(name='Arial', size=10, color='1155CC', underline='single')
THIN_BORDER = Border(
    bottom=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='E0E0E0')
)
ALT_FILL = PatternFill('solid', fgColor='F7F9FC')
HIGH_FILL = PatternFill('solid', fgColor='E8F5E9')
MED_FILL = PatternFill('solid', fgColor='FFF8E1')
LOW_FILL = PatternFill('solid', fgColor='FFEBEE')

COL_WIDTHS = {
    'Page URL': 45, 'Page Type': 16, 'HTTP Status Code': 12,
    'Page Title': 50, 'Meta Description': 50,
    'Primary Keyword': 35, 'Secondary Keywords': 45,
    'Total Keywords Ranking': 14, 'KW Clicks': 12, 'KW Impressions': 14,
    'KW Avg Position': 12, 'GSC Clicks': 12, 'GSC Impressions': 14,
    'GSC CTR': 10, 'GSC Avg Pos.': 12, 'PA': 8, 'Inbound Links': 14,
    'Linking Domains': 14, 'Priority': 10, 'Optimization Status': 16,
}

TECHNICAL_SEGS = {
    'Crawl Depth > 1', 'Crawl Depth > 2', 'Crawl Depth > 3',
    'Crawl Depth > 4', 'Crawl Depth > 5',
    'Errors - 3XX Redirects', 'Errors - 4XX', 'Errors - 5XX',
    '301 w/ 5+ inlinks', '4xx error 1 w/ 5+ inlinks_1', 'Top 500',
}

# Patterns for classifying page types from URL paths
URL_CLASSIFIERS = [
    (r'/blog/', 'Blog'),
    (r'/help/', 'Help/Support'),
    (r'/support/', 'Help/Support'),
    (r'/learn/', 'Learning Center'),
    (r'/resources/', 'Learning Center'),
    (r'/education/', 'Learning Center'),
    (r'/products/', 'Product Page'),
    (r'/features/', 'Product Page'),
    (r'/pricing', 'Product Page'),
    (r'/tools/', 'Freemium Tool'),
    (r'/checker', 'Freemium Tool'),
    (r'/analyzer', 'Freemium Tool'),
    (r'/guide', 'Guide'),
    (r'/how-to/', 'Guide'),
    (r'/about/', 'About'),
    (r'/team/', 'About'),
    (r'/careers', 'About'),
    (r'/community/', 'Community'),
    (r'/forum/', 'Community'),
    (r'/users/', 'Community'),
    (r'/webinar', 'Webinar'),
    (r'/case-stud', 'Case Study'),
    (r'/white-paper', 'White Paper'),
    (r'/press/', 'Press'),
    (r'/login', 'Login'),
    (r'/api/', 'API Docs'),
]

PRIORITY_DEFAULTS = {
    'Home Page': 'High',
    'Product Page': 'High',
    'Freemium Tool': 'High',
    'Learning Center': 'Medium',
    'Help/Support': 'Medium',
    'Guide': 'Medium',
    'Blog': 'Medium',
    'Case Study': 'Medium',
    'Webinar': 'Medium',
    'White Paper': 'Medium',
    'Community': 'Low',
    'About': 'Low',
    'Press': 'Low',
    'Login': 'Low',
    'API Docs': 'Low',
    'Event': 'Low',
    'Additional Content': 'Low',
    "What's New": 'Low',
    'All Domain Analysis': 'High',
    'Other': 'Low',
}


def fuzzy_col(df, candidates):
    """Find the first matching column name from a list of candidates (case-insensitive)."""
    lower_cols = {c.lower().strip(): c for c in df.columns}
    for cand in candidates:
        if cand.lower().strip() in lower_cols:
            return lower_cols[cand.lower().strip()]
    return None


def detect_columns(crawl_df, ranking_df):
    """Map source columns to workbook columns, adapting to whatever the files contain."""
    crawl_map = {}
    col_candidates = {
        'url': ['Address', 'URL', 'Page URL', 'Page', 'Landing Page'],
        'status': ['Status Code', 'Status', 'HTTP Status', 'HTTP Status Code'],
        'title': ['Title 1', 'Title', 'Page Title', 'SEO Title'],
        'segments': ['Segments', 'Segment', 'Category', 'Content Type', 'Page Type'],
        'meta_desc': ['Meta Description 1', 'Meta Description', 'Description'],
        'clicks': ['Clicks', 'GSC Clicks', 'Search Clicks'],
        'impressions': ['Impressions', 'GSC Impressions', 'Search Impressions'],
        'ctr': ['CTR', 'GSC CTR', 'Click Through Rate'],
        'position': ['Position', 'Average Position', 'GSC Avg Pos.', 'Avg Position'],
        'pa': ['Moz Page Authority', 'Page Authority', 'PA'],
        'external_links': ['Moz External Pages to Page', 'External Links', 'Inbound Links', 'Backlinks'],
        'root_domains': ['Moz Root Domains to Page', 'Root Domains', 'Linking Domains', 'Referring Domains'],
        'embeddings': ['Extract embeddings from page content', 'Embeddings', 'Embedding', 'Content Embedding'],
    }
    for key, cands in col_candidates.items():
        found = fuzzy_col(crawl_df, cands)
        if found:
            crawl_map[key] = found

    rank_map = {}
    rank_candidates = {
        'query': ['Query', 'Keyword', 'Search Query', 'Term'],
        'url': ['Landing Page', 'URL', 'Page', 'Address'],
        'clicks': ['Url Clicks', 'Clicks', 'URL Clicks'],
        'impressions': ['Impressions', 'URL Impressions', 'Url Impressions'],
        'position': ['Average Position', 'Position', 'Avg Position', 'URL Position'],
    }
    for key, cands in rank_candidates.items():
        found = fuzzy_col(ranking_df, cands)
        if found:
            rank_map[key] = found

    return crawl_map, rank_map


def clean_segment(seg):
    if pd.isna(seg):
        return ''
    parts = [p.strip() for p in str(seg).split(',')]
    # Remove technical segments and crawl depth variants
    cleaned = [p for p in parts if p not in TECHNICAL_SEGS and not re.match(r'Crawl Depth', p, re.I)]
    return ', '.join(cleaned) if cleaned else ''


def classify_page_type(seg_clean, url, domain):
    """Classify a page into a content type using cleaned segment + URL patterns."""
    url_lower = str(url).lower() if pd.notna(url) else ''

    # Homepage detection
    url_path = url_lower.replace(f'https://{domain}', '').replace(f'http://{domain}', '').rstrip('/')
    if url_path == '' or url_path == '/':
        return 'Home Page'

    # Use segment if it's meaningful
    if seg_clean:
        seg_lower = seg_clean.lower()
        for keyword, ptype in [
            ('blog', 'Blog'), ('help', 'Help/Support'), ('support', 'Help/Support'),
            ('learn', 'Learning Center'), ('resource', 'Learning Center'),
            ('product', 'Product Page'), ('freemium', 'Freemium Tool'),
            ('guide', 'Guide'), ('case stud', 'Case Study'), ('webinar', 'Webinar'),
            ('white paper', 'White Paper'), ('about', 'About'), ('press', 'Press'),
            ('community', 'Community'), ('login', 'Login'), ('mozcon', 'Event'),
        ]:
            if keyword in seg_lower:
                return ptype

    # Fall back to URL patterns
    for pattern, ptype in URL_CLASSIFIERS:
        if re.search(pattern, url_lower):
            return ptype

    return seg_clean if seg_clean else 'Other'


def extract_domain(urls):
    """Extract the most common domain from a series of URLs."""
    from urllib.parse import urlparse
    domains = urls.dropna().apply(lambda u: urlparse(str(u)).netloc).value_counts()
    return domains.index[0] if len(domains) > 0 else ''


def assign_keywords(ranking_df, rank_map):
    """Score and assign primary/secondary keywords per URL from ranking data."""
    rdf = ranking_df.dropna(subset=[rank_map['query'], rank_map['url']]).copy()
    rdf['_clicks'] = pd.to_numeric(rdf[rank_map['clicks']], errors='coerce').fillna(0)
    rdf['_impressions'] = pd.to_numeric(rdf[rank_map.get('impressions', '_missing')], errors='coerce').fillna(0) if 'impressions' in rank_map else 0
    rdf['_position'] = pd.to_numeric(rdf[rank_map.get('position', '_missing')], errors='coerce').fillna(100) if 'position' in rank_map else 50

    rdf['_score'] = rdf['_clicks'] * 3 + rdf['_impressions'] * 0.1 + (100 - rdf['_position'].clip(1, 100)) * 5
    rdf = rdf.sort_values('_score', ascending=False)

    # Primary keyword (top 1)
    top1 = rdf.groupby(rank_map['url']).first().reset_index()
    top1 = top1[[rank_map['url'], rank_map['query'], '_clicks', '_impressions', '_position']]
    top1.columns = ['_url', 'Primary Keyword', 'KW Clicks', 'KW Impressions', 'KW Avg Position']

    # Secondary keywords (2nd-4th)
    def get_secondary(g):
        return ', '.join(g[rank_map['query']].iloc[1:4].tolist()) if len(g) > 1 else ''
    top23 = rdf.groupby(rank_map['url']).apply(get_secondary, include_groups=False).reset_index()
    top23.columns = ['_url', 'Secondary Keywords']

    # Keyword count
    kw_count = rdf.groupby(rank_map['url'])[rank_map['query']].count().reset_index()
    kw_count.columns = ['_url', 'Total Keywords Ranking']

    return top1, top23, kw_count


def title_fallback_keyword(title):
    """Extract a keyword phrase from page title by stripping site name suffix."""
    if pd.isna(title):
        return ''
    cleaned = re.sub(r'\s*[-–—|]\s*[\w\s.]+$', '', str(title)).strip()
    return cleaned[:80] if cleaned else ''


def write_data_sheet(ws, data, display_cols, tab_color):
    """Write a formatted data sheet with headers, formatting, and validation."""
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = 'A2'

    cols = [c for c in display_cols if c in data.columns]

    for ci, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(col, 12)

    for ri, (_, row) in enumerate(data.iterrows(), 2):
        for ci, col in enumerate(cols, 1):
            val = row.get(col, '')
            if pd.isna(val):
                val = ''
            elif isinstance(val, float) and val == int(val) and col not in ('GSC CTR', 'KW Avg Position', 'GSC Avg Pos.'):
                val = int(val)

            # Escape formula-like strings
            if isinstance(val, str) and len(val) > 0 and val[0] in ('=', '+', '-', '@'):
                val = "'" + val

            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical='center', wrap_text=(col in ('Page Title', 'Meta Description', 'Secondary Keywords')))

            if col == 'Page URL' and isinstance(val, str) and val.startswith('http'):
                cell.font = LINK_FONT
            if col == 'GSC CTR' and isinstance(val, (int, float)):
                cell.number_format = '0.00%'
            elif col in ('KW Avg Position', 'GSC Avg Pos.') and isinstance(val, (int, float)):
                cell.number_format = '0.0'
            elif col in ('KW Clicks', 'KW Impressions', 'GSC Clicks', 'GSC Impressions', 'Inbound Links', 'Linking Domains', 'Total Keywords Ranking'):
                cell.number_format = '#,##0'
            if col == 'Priority':
                if val == 'High': cell.fill = HIGH_FILL
                elif val == 'Medium': cell.fill = MED_FILL
                elif val == 'Low': cell.fill = LOW_FILL
            elif ri % 2 == 0:
                cell.fill = ALT_FILL

    # Data validation for Optimization Status
    if 'Optimization Status' in cols:
        idx = cols.index('Optimization Status') + 1
        dv = DataValidation(type='list', formula1='"Backlog,To Optimize,In Progress,Optimized"', allow_blank=True)
        ws.add_data_validation(dv)
        cl = get_column_letter(idx)
        dv.add(f'{cl}2:{cl}{len(data)+1}')

    # Data validation for Priority
    if 'Priority' in cols:
        idx = cols.index('Priority') + 1
        dv = DataValidation(type='list', formula1='"High,Medium,Low"', allow_blank=True)
        ws.add_data_validation(dv)
        cl = get_column_letter(idx)
        dv.add(f'{cl}2:{cl}{len(data)+1}')

    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(data)+1}'


def build_workbook(crawl_path, ranking_path, output_path, site_name=None):
    """Main entry point: build the SEO workbook."""
    # Load data
    crawl_df = pd.read_excel(crawl_path)
    if str(ranking_path).endswith('.csv'):
        ranking_df = pd.read_csv(ranking_path)
    else:
        ranking_df = pd.read_excel(ranking_path)

    print(f"Crawl data: {crawl_df.shape[0]} rows, {crawl_df.shape[1]} columns")
    print(f"Ranking data: {ranking_df.shape[0]} rows, {ranking_df.shape[1]} columns")

    # Detect columns
    crawl_map, rank_map = detect_columns(crawl_df, ranking_df)
    print(f"\nCrawl columns mapped: {json.dumps(crawl_map, indent=2)}")
    print(f"Ranking columns mapped: {json.dumps(rank_map, indent=2)}")

    if 'url' not in crawl_map:
        sys.exit("ERROR: Could not find a URL column in the crawl export.")
    if 'query' not in rank_map or 'url' not in rank_map:
        sys.exit("ERROR: Could not find query and URL columns in the ranking data.")

    # Detect domain
    domain = extract_domain(crawl_df[crawl_map['url']])
    if not site_name:
        site_name = domain
    print(f"Domain: {domain}")

    # Clean segments and classify page types
    if 'segments' in crawl_map:
        crawl_df['_seg_clean'] = crawl_df[crawl_map['segments']].apply(clean_segment)
    else:
        crawl_df['_seg_clean'] = ''

    crawl_df['Page Type'] = crawl_df.apply(
        lambda r: classify_page_type(r['_seg_clean'], r[crawl_map['url']], domain), axis=1
    )

    print(f"\nPage Type Distribution:")
    for pt, count in crawl_df['Page Type'].value_counts().items():
        print(f"  {pt}: {count}")

    # Build output dataframe
    col_mapping = {
        crawl_map['url']: 'Page URL',
        'Page Type': 'Page Type',
    }
    if 'status' in crawl_map: col_mapping[crawl_map['status']] = 'HTTP Status Code'
    if 'title' in crawl_map: col_mapping[crawl_map['title']] = 'Page Title'
    if 'meta_desc' in crawl_map: col_mapping[crawl_map['meta_desc']] = 'Meta Description'
    if 'clicks' in crawl_map: col_mapping[crawl_map['clicks']] = 'GSC Clicks'
    if 'impressions' in crawl_map: col_mapping[crawl_map['impressions']] = 'GSC Impressions'
    if 'ctr' in crawl_map: col_mapping[crawl_map['ctr']] = 'GSC CTR'
    if 'position' in crawl_map: col_mapping[crawl_map['position']] = 'GSC Avg Pos.'
    if 'pa' in crawl_map: col_mapping[crawl_map['pa']] = 'PA'
    if 'external_links' in crawl_map: col_mapping[crawl_map['external_links']] = 'Inbound Links'
    if 'root_domains' in crawl_map: col_mapping[crawl_map['root_domains']] = 'Linking Domains'
    if 'embeddings' in crawl_map: col_mapping[crawl_map['embeddings']] = 'Embedding'

    source_cols = [c for c in col_mapping.keys() if c in crawl_df.columns]
    df = crawl_df[source_cols].copy()
    df = df.rename(columns=col_mapping)

    # Keyword assignment
    top1, top23, kw_count = assign_keywords(ranking_df, rank_map)
    df = df.merge(top1, left_on='Page URL', right_on='_url', how='left').drop(columns=['_url'], errors='ignore')
    df = df.merge(top23, left_on='Page URL', right_on='_url', how='left').drop(columns=['_url'], errors='ignore')
    df = df.merge(kw_count, left_on='Page URL', right_on='_url', how='left').drop(columns=['_url'], errors='ignore')

    # Title fallback for missing keywords
    if 'Page Title' in df.columns:
        mask = df['Primary Keyword'].isna()
        df.loc[mask, 'Primary Keyword'] = df.loc[mask, 'Page Title'].apply(title_fallback_keyword)

    # Priority
    df['Priority'] = df['Page Type'].map(PRIORITY_DEFAULTS).fillna('Low')

    # Optimization Status
    df['Optimization Status'] = 'Backlog'

    # Column order
    display_cols = [
        'Page URL', 'Page Type', 'HTTP Status Code', 'Page Title', 'Meta Description',
        'Primary Keyword', 'Secondary Keywords', 'Total Keywords Ranking',
        'KW Clicks', 'KW Impressions', 'KW Avg Position',
        'GSC Clicks', 'GSC Impressions', 'GSC CTR', 'GSC Avg Pos.',
        'PA', 'Inbound Links', 'Linking Domains',
        'Priority', 'Optimization Status',
    ]
    existing_display = [c for c in display_cols if c in df.columns]

    # Sort by priority then clicks
    priority_sort = {'High': 0, 'Medium': 1, 'Low': 2}
    df['_psort'] = df['Priority'].map(priority_sort)
    sort_cols = ['_psort']
    if 'GSC Clicks' in df.columns:
        sort_cols.append('GSC Clicks')
        df = df.sort_values(sort_cols, ascending=[True, False])
    else:
        df = df.sort_values(sort_cols)
    df = df.drop(columns=['_psort'])

    # Sanitize any #NAME?-like strings
    for col in df.select_dtypes(include='object').columns:
        df[col] = df[col].apply(lambda v: str(v).replace('#NAME?', '#name') if isinstance(v, str) and '#NAME?' in v else v)

    # ============================================================
    # BUILD WORKBOOK
    # ============================================================
    wb = Workbook()

    # --- Summary sheet ---
    ws_sum = wb.active
    ws_sum.title = 'Summary'
    ws_sum.sheet_properties.tabColor = '1B2A4A'

    summary_rows = [
        [f'{site_name} SEO Workbook', '', '', ''],
        ['', '', '', ''],
        ['Metric', 'Count', '', ''],
        ['Total Pages Crawled', len(df), '', ''],
        ['Pages with Ranking Keywords', int((df['Total Keywords Ranking'].notna() & (df['Total Keywords Ranking'] > 0)).sum()) if 'Total Keywords Ranking' in df.columns else 0, '', ''],
        ['High Priority Pages', int((df['Priority'] == 'High').sum()), '', ''],
        ['Medium Priority Pages', int((df['Priority'] == 'Medium').sum()), '', ''],
        ['Low Priority Pages', int((df['Priority'] == 'Low').sum()), '', ''],
        ['', '', '', ''],
        ['Page Type', 'Count', 'GSC Clicks', 'Avg PA'],
    ]
    for row in summary_rows:
        ws_sum.append(row)

    for pt in df['Page Type'].value_counts().index:
        sub = df[df['Page Type'] == pt]
        clicks_total = int(sub['GSC Clicks'].sum()) if 'GSC Clicks' in sub.columns and sub['GSC Clicks'].notna().any() else ''
        avg_pa = round(sub['PA'].mean(), 1) if 'PA' in sub.columns and sub['PA'].notna().any() else ''
        ws_sum.append([pt, len(sub), clicks_total, avg_pa])

    ws_sum['A1'].font = Font(name='Arial', bold=True, size=18, color='1B2A4A')
    for cell in ws_sum[3]:
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
    for cell in ws_sum[10]:
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
    ws_sum.column_dimensions['A'].width = 30
    ws_sum.column_dimensions['B'].width = 15
    ws_sum.column_dimensions['C'].width = 15
    ws_sum.column_dimensions['D'].width = 12

    # --- Per-page-type sheets ---
    TAB_COLORS = ['4CAF50', '2196F3', '9C27B0', 'FF9800', '00BCD4', '795548', 'F44336',
                  'E91E63', '3F51B5', '009688', 'FFC107', '8BC34A']
    page_types_ordered = df['Page Type'].value_counts().index.tolist()
    # Put high-priority types first
    high_types = [pt for pt in page_types_ordered if PRIORITY_DEFAULTS.get(pt) == 'High']
    med_types = [pt for pt in page_types_ordered if PRIORITY_DEFAULTS.get(pt) == 'Medium']
    low_types = [pt for pt in page_types_ordered if PRIORITY_DEFAULTS.get(pt, 'Low') == 'Low']
    ordered = high_types + med_types + low_types

    for i, pt in enumerate(ordered):
        sub = df[df['Page Type'] == pt].copy()
        if 'GSC Clicks' in sub.columns:
            sub = sub.sort_values('GSC Clicks', ascending=False)
        sheet_name = pt[:31]  # Excel 31-char limit
        # Strip unicode variation selectors and invalid Excel sheet name chars
        sheet_name = sheet_name.replace('\ufe0f', '')
        for ch in ['/', '\\', '?', '*', '[', ']', ':']:
            sheet_name = sheet_name.replace(ch, '-')
        color = TAB_COLORS[i % len(TAB_COLORS)]
        ws = wb.create_sheet(sheet_name)
        write_data_sheet(ws, sub, existing_display, color)
        print(f"  Sheet '{sheet_name}': {len(sub)} rows")

    # --- All Pages sheet ---
    ws_all = wb.create_sheet('All Pages')
    all_sorted = df.copy()
    if 'GSC Clicks' in all_sorted.columns:
        all_sorted = all_sorted.sort_values('GSC Clicks', ascending=False)
    write_data_sheet(ws_all, all_sorted, existing_display, '607D8B')
    print(f"  Sheet 'All Pages': {len(all_sorted)} rows")

    # --- Embeddings sheet (if data exists) ---
    if 'Embedding' in df.columns:
        emb_df = df[['Page URL', 'Page Type', 'Primary Keyword', 'Embedding']].copy()
        emb_df = emb_df[emb_df['Embedding'].notna() & (emb_df['Embedding'] != '')]
        if len(emb_df) > 0:
            ws_emb = wb.create_sheet('Embeddings')
            ws_emb.sheet_properties.tabColor = '9E9E9E'
            for ci, h in enumerate(['Page URL', 'Page Type', 'Primary Keyword', 'Embedding'], 1):
                cell = ws_emb.cell(row=1, column=ci, value=h)
                cell.font = HEADER_FONT; cell.fill = HEADER_FILL
            ws_emb.column_dimensions['A'].width = 45
            ws_emb.column_dimensions['B'].width = 16
            ws_emb.column_dimensions['C'].width = 35
            ws_emb.column_dimensions['D'].width = 50
            ws_emb.freeze_panes = 'A2'
            for ri, (_, row) in enumerate(emb_df.iterrows(), 2):
                for ci, col in enumerate(['Page URL', 'Page Type', 'Primary Keyword', 'Embedding'], 1):
                    val = row[col] if pd.notna(row[col]) else ''
                    if col == 'Embedding' and isinstance(val, str) and len(val) > 200:
                        val = val[:200] + '...'
                    ws_emb.cell(row=ri, column=ci, value=val).font = DATA_FONT
            print(f"  Sheet 'Embeddings': {len(emb_df)} rows")

    # Remove Embedding from main display columns (it's in its own sheet)
    # (already excluded from existing_display since we didn't add it)

    # Save
    wb.save(output_path)
    print(f"\nWorkbook saved to: {output_path}")
    print(f"Total sheets: {len(wb.sheetnames)}")
    return output_path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Build SEO Workbook from crawl + ranking data')
    parser.add_argument('--crawl', required=True, help='Path to Screaming Frog internal URLs export (.xlsx)')
    parser.add_argument('--ranking', required=True, help='Path to ranking data (.csv or .xlsx)')
    parser.add_argument('--output', required=True, help='Output path for the workbook (.xlsx)')
    parser.add_argument('--site-name', default=None, help='Site name for the workbook title')
    args = parser.parse_args()
    build_workbook(args.crawl, args.ranking, args.output, args.site_name)
