#!/usr/bin/env python3
"""Convert the Markdown content package to a clean Word document.

Usage:
  python md_to_docx.py package.md package.docx

Fallback for environments without a docx skill. Requires python-docx
(pip install python-docx). Handles headings, paragraphs, bold, italics, inline code,
links, bullet and numbered lists, blockquotes (used for [MEDIA-xx] markers), pipe
tables, fenced code blocks (JSON-LD), and horizontal rules.
"""
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
FLAG = re.compile(r"\[(EXPERT INPUT|VERIFY|PLACEHOLDER)[^\]]*\]")


def add_link(par, text, url):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FBF")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_text(par, text):
    """Add text, highlighting [EXPERT INPUT]/[VERIFY]/[PLACEHOLDER] flags for editors."""
    pos = 0
    for m in FLAG.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        r = par.add_run(m.group(0))
        r.bold = True
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def inline(par, text):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = par.add_run(chunk[1:-1])
            r.font.name = "Consolas"
        elif chunk.startswith("[") and "](" in chunk and chunk.endswith(")"):
            label, url = chunk[1:-1].split("](", 1)
            if label.split(":")[0] in ("EXPERT INPUT", "VERIFY", "PLACEHOLDER"):
                add_text(par, chunk)
            else:
                add_link(par, label, url)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        else:
            add_text(par, chunk)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=width)
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(width):
            p = t.cell(i, j).paragraphs[0]
            inline(p, r[j] if j < len(r) else "")
            if i == 0:
                for run in p.runs:
                    run.bold = True


def convert(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            continue
        elif m := re.match(r"(#{1,6})\s+(.*)", s):
            h = doc.add_heading(level=min(len(m.group(1)), 4))
            inline(h, m.group(2))
        elif re.fullmatch(r"-{3,}|\*{3,}", s):
            doc.add_paragraph("_" * 40)
        elif s.startswith(">"):
            p = doc.add_paragraph()
            inline(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif m := re.match(r"\s*[-*]\s+(\[[ x]\]\s+)?(.*)", line):
            inline(doc.add_paragraph(style="List Bullet"), (m.group(1) or "") + m.group(2))
        elif m := re.match(r"\s*\d+[.)]\s+(.*)", line):
            inline(doc.add_paragraph(style="List Number"), m.group(1))
        elif s:
            inline(doc.add_paragraph(), s)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    convert(sys.argv[1], sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
