"""Offline regression suite: python scripts/test_integrity.py."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
import build_workbook as b
from integrity import *
from reporting import assert_outputs
from schema_validate import validate

BASE=json.loads(Path(__file__).with_name('example_findings.json').read_text())

class IntegrityTests(unittest.TestCase):
    def fixture(self):
        d=copy.deepcopy(BASE)
        prepare(d)
        return d

    def test_overlap_unknown(self):
        self.assertIsNone(direct_follow_complement(dict(domains_to_page=54,indirect_domains=53,nofollow_domains=21)))

    def test_overlap_known(self):
        r=dict(domains_to_page=54,indirect_domains=53,nofollow_domains=21,indirect_nofollow_overlap=20,composition_evidence='same sets')
        self.assertEqual(direct_follow_complement(r),0)
        r['indirect_nofollow_overlap']=0
        with self.assertRaises(ValueError): direct_follow_complement(r)

    def test_normalize_client_labels(self):
        for s in ('CLIENT (moz.com)','https://WWW.MOZ.COM/products/local','moz.com.'):
            self.assertEqual(normalize_domain(s),'moz.com')
        self.assertEqual(normalize_domain('https://publisher.blogspot.com/a'),'publisher.blogspot.com')
        self.assertEqual(normalize_domain('example.co.uk/path'),'example.co.uk')
        verdict=b._seller_verdict(['CLIENT (moz.com)'],'moz.com/products/local',6)
        self.assertIn('hit the client',verdict)
        self.assertNotIn('bought',b._seller_verdict(['competitor.example'],'client.example',6))

    def test_partial_competitors(self):
        d=self.fixture()
        d['competitors'].append({'domain':'missing.example'})
        self.assertTrue(b.check_coverage(d))
        with self.assertRaises(ValueError): validate_partial_reasons(d)

    def test_reasoned_partial(self):
        d=self.fixture()
        c=d['coverage']['recently_earned']
        c['completed']=[]
        c['omitted']={'competitor.example':'Quota cutoff'}
        self.assertTrue(b.check_coverage(d))
        validate_partial_reasons(d)

    def test_unverified_high_score(self):
        r=self.fixture()['scored_targets'][-1]
        self.assertEqual(decision(r),'Verify first')

    def test_traffic_proxy_required(self):
        r=self.fixture()['scored_targets'][0]
        del r['traffic_source']
        self.assertEqual(decision(r),'Verify first')

    def test_imported_score_not_auditable(self):
        r=self.fixture()['scored_targets'][0]
        for k in FACTORS: del r[k]
        r['score']=100
        self.assertEqual(decision(r),'Verify first')

    def test_counts_domains_distinct(self):
        c=counts(self.fixture())
        self.assertEqual((c['opportunities'],c['unique_domains'],c['Pursue'],c['Backlog'],c['Discard'],c['Verify first']),(4,3,1,1,1,1))

    def test_override_reason(self):
        d=self.fixture()
        d['scored_targets'][1]['execution_decision']='Hygiene quick win'
        with self.assertRaises(ValueError): prepare(d)
        d['scored_targets'][1]['override_reason']='Verified free listing correction'
        prepare(d)

    def test_unverified_override_blocked(self):
        d=self.fixture()
        d['scored_targets'][-1].update(execution_decision='Pursue now',override_reason='High score')
        with self.assertRaises(ValueError): prepare(d)

    def test_plan_must_reference_scorecard(self):
        d=self.fixture()
        d['outreach_plan']=[{'opportunity_id':'unknown'}]
        with self.assertRaises(ValueError): validate_plan(d)

    def test_approved_proof_only(self):
        d=self.fixture()
        b.derive_outreach_plan(d)
        d['outreach_plan'][0]['proof_point_ids']=['invented']
        with self.assertRaises(ValueError): validate_plan(d)

    def test_live_serp_provenance(self):
        d=self.fixture()
        d['live_serp']={'supplied':True}
        with self.assertRaises(ValueError): prepare(d)

    def test_schema_factor_bounds(self):
        schema=json.loads(Path(__file__).with_name('findings.schema.json').read_text())
        validate(BASE,schema)
        d=copy.deepcopy(BASE)
        d['scored_targets'][0]['relevance']=31
        with self.assertRaises(ValueError): validate(d,schema)

    def test_build_reconcile_and_detect_wrong_column(self):
        d=self.fixture()
        b.check_coverage(d)
        b.derive_outreach_plan(d)
        validate_plan(d)
        with tempfile.TemporaryDirectory() as temp:
            x=str(Path(temp)/'sample.xlsx'); m=str(Path(temp)/'summary.md')
            b.build(d,x); b.write_summary(d,m,'sample.xlsx')
            assert_outputs(d,x,m,b.SCORED_COLS)
            cached=load_workbook(x,data_only=True)
            summary={r[0]:r[1] for r in cached['Summary'].iter_rows(values_only=True) if r[0]}
            self.assertEqual(summary['Pursue'],1)
            self.assertEqual(summary['Verify first'],1)
            wb=load_workbook(x)
            for row in wb['Summary']:
                if row[0].value=='Pursue':
                    row[1].value=row[1].value.replace('!P2:P5','!O2:O5')
            wb.save(x)
            with self.assertRaises(AssertionError): assert_outputs(d,x,m,b.SCORED_COLS)

    def test_empty_build(self):
        d=self.fixture(); d['scored_targets']=[]
        with tempfile.TemporaryDirectory() as temp:
            x=str(Path(temp)/'empty.xlsx'); m=str(Path(temp)/'empty.md')
            b.build(d,x); b.write_summary(d,m,'empty.xlsx'); assert_outputs(d,x,m,b.SCORED_COLS)

if __name__=='__main__': unittest.main(verbosity=2)
